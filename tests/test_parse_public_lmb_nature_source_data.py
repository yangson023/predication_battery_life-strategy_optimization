import csv
import json
import tempfile
import unittest
from pathlib import Path

import openpyxl

from modules.data_pipeline.parse_public_lmb_nature_source_data import (
    parse_public_lmb_nature_source_data,
)


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def build_synthetic_workbook(path: Path, workbook_kind: str = "63303") -> None:
    workbook = openpyxl.Workbook()
    default = workbook.active
    workbook.remove(default)

    cycle = workbook.create_sheet("Figure 3a" if workbook_kind == "63303" else "Fig. 5c")
    if workbook_kind == "63303":
        cycle.append(["Cycle number", "6-3-cell1-Cap", "6-3-cell1-CE", "6-3-cell2-Cap", "6-3-cell2-CE"])
        cycle.append([1, 147.2, 94.0, 146.3, 94.1])
        cycle.append([2, 146.8, 99.3, 146.0, 99.2])
    else:
        cycle.append(["", "CC1", "", "MPC1", ""])
        cycle.append(["Cycle", "Capacity", "CE", "Capacity", "CE"])
        cycle.append([1, 2.0, 67.0, 2.2, 73.0])
        cycle.append([2, 2.9, 99.5, 3.0, 101.4])

    time = workbook.create_sheet("Figure S7")
    time.append(["3-1-cell1", "", "3-1-cell2", ""])
    time.append(["Time (hours)", "Voltage (V)", "Time (hours)", "Voltage (V)"])
    time.append([0, 0.19, 0, 0.17])
    time.append([0.00278, 1.02, 0.00278, 1.03])

    descriptor = workbook.create_sheet("Figure S16")
    descriptor.append(["", "ELi (V)", "Raman peak center (cm-1)", "Discharge capacity at 20th cycle"])
    descriptor.append(["6-3", -3.18, 745.8, 124.1])
    descriptor.append(["7-1", -3.22, 731.8, 115.6])

    unknown = workbook.create_sheet("Mystery")
    unknown.append(["foo", "bar"])
    unknown.append([1, 2])

    workbook.save(path)


class PublicLmbNatureSourceDataParserTests(unittest.TestCase):
    def test_reads_multi_sheet_workbook_and_writes_inventory(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            workbook_path = root / "41467_2025_63303_MOESM3_ESM.xlsx"
            build_synthetic_workbook(workbook_path, "63303")

            parse_public_lmb_nature_source_data(
                input_files=[workbook_path],
                output_root=root / "out",
                overwrite=True,
                max_preview_rows=10,
            )

            workbooks = read_csv(root / "out" / "source_workbook_inventory.csv")
            sheets = read_csv(root / "out" / "source_sheet_inventory.csv")
            self.assertEqual(workbooks[0]["sheet_count"], "4")
            self.assertEqual(len(sheets), 4)
            self.assertEqual({row["model_training_allowed"] for row in sheets}, {"False"})

    def test_cycle_capacity_ce_sheet_parses_to_long_table(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            workbook_path = root / "41467_2025_63303_MOESM3_ESM.xlsx"
            build_synthetic_workbook(workbook_path, "63303")

            parse_public_lmb_nature_source_data(
                input_files=[workbook_path],
                output_root=root / "out",
                overwrite=True,
                max_preview_rows=10,
                selected_sheets=["Figure 3a"],
            )

            rows = read_csv(root / "out" / "public_lmb_source_long_preview.csv")
            self.assertTrue(any(row["capacity"] for row in rows))
            self.assertTrue(any(row["CE"] for row in rows))
            self.assertEqual({row["feature_family"] for row in rows}, {"cycle_capacity_ce"})
            self.assertEqual({row["model_training_allowed"] for row in rows}, {"False"})

    def test_time_voltage_sheet_parses_to_long_table_and_is_not_full_cell_for_licu_mechanism(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            workbook_path = root / "41467_2025_63303_MOESM3_ESM.xlsx"
            build_synthetic_workbook(workbook_path, "63303")

            parse_public_lmb_nature_source_data(
                input_files=[workbook_path],
                output_root=root / "out",
                overwrite=True,
                max_preview_rows=10,
                selected_sheets=["Figure S7"],
            )

            rows = read_csv(root / "out" / "public_lmb_source_long_preview.csv")
            self.assertTrue(rows)
            self.assertEqual({row["feature_family"] for row in rows}, {"time_voltage_current"})
            self.assertEqual({row["cell_scope"] for row in rows}, {"lmb_mechanism_test_not_full_cell"})

    def test_unknown_sheet_generates_warning_when_selected(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            workbook_path = root / "41467_2025_63303_MOESM3_ESM.xlsx"
            build_synthetic_workbook(workbook_path, "63303")

            parse_public_lmb_nature_source_data(
                input_files=[workbook_path],
                output_root=root / "out",
                overwrite=True,
                max_preview_rows=10,
                selected_sheets=["Mystery"],
            )

            warnings = read_csv(root / "out" / "source_data_parser_warnings.csv")
            self.assertTrue(warnings)
            self.assertIn("unsupported", warnings[0]["warning_message"])

    def test_c20_proxy_is_not_full_rul_or_eol(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            workbook_path = root / "41467_2025_63303_MOESM3_ESM.xlsx"
            build_synthetic_workbook(workbook_path, "63303")

            parse_public_lmb_nature_source_data(
                input_files=[workbook_path],
                output_root=root / "out",
                overwrite=True,
                max_preview_rows=10,
                selected_sheets=["Figure S16"],
            )

            rows = read_csv(root / "out" / "public_lmb_source_long_preview.csv")
            report = json.loads((root / "out" / "public_lmb_source_data_audit_report.json").read_text(encoding="utf-8"))
            self.assertTrue(rows)
            self.assertTrue(any("C20_proxy_not_RUL" in row["feature_family"] for row in rows))
            self.assertFalse(report["model_training_allowed"])

    def test_report_has_no_model_performance_claim(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            workbook_path = root / "41467_2025_66271_MOESM3_ESM.xlsx"
            build_synthetic_workbook(workbook_path, "66271")

            parse_public_lmb_nature_source_data(
                input_files=[workbook_path],
                output_root=root / "out",
                overwrite=True,
                max_preview_rows=10,
                selected_sheets=["Fig. 5c"],
            )

            report_text = (root / "out" / "public_lmb_source_data_audit_report.md").read_text(encoding="utf-8").lower()
            report = json.loads((root / "out" / "public_lmb_source_data_audit_report.json").read_text(encoding="utf-8"))
            self.assertFalse(report["model_training_allowed"])
            self.assertFalse(report["formal_model_performance_claimed"])
            for blocked in ["auc", "f1", "rmse", "r2"]:
                self.assertNotIn(blocked, report_text)


if __name__ == "__main__":
    unittest.main()
