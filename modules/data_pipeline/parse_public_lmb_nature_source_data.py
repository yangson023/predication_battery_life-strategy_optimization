"""Parse public LMB Nature source-data workbooks for audit-only use.

This parser creates workbook/sheet inventories and a tiny long-format preview
from selected Nature Communications source-data Excel sheets. It does not train
models, create processed battery datasets, generate labels, or enter the RUL
prediction pipeline.
"""

from __future__ import annotations

import argparse
import csv
import json
import re
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import openpyxl


SHEET_TYPES = {
    "cycle_capacity_ce",
    "time_voltage_current",
    "capacity_voltage_curve",
    "electrolyte_descriptor",
    "spectroscopy_or_characterization",
    "active_learning_metric",
    "figure_source_unknown",
}

WORKBOOK_INVENTORY_COLUMNS = [
    "workbook_name",
    "source_path",
    "paper_id",
    "sheet_count",
    "source_data_audit_only",
    "model_training_allowed",
]

SHEET_INVENTORY_COLUMNS = [
    "workbook_name",
    "paper_id",
    "sheet_name",
    "n_rows",
    "n_cols",
    "header_preview",
    "inferred_sheet_type",
    "dataset_role",
    "cell_scope",
    "training_allowed_now",
    "audit_use",
    "source_data_audit_only",
    "model_training_allowed",
]

PARSE_MANIFEST_COLUMNS = [
    "workbook_name",
    "paper_id",
    "sheet_name",
    "inferred_sheet_type",
    "parse_status",
    "parsed_rows",
    "parser_warning",
    "source_data_audit_only",
    "model_training_allowed",
]

LONG_COLUMNS = [
    "workbook_name",
    "paper_id",
    "sheet_name",
    "trace_id",
    "inferred_cell_or_condition",
    "cycle_index",
    "time_index",
    "capacity",
    "normalized_capacity",
    "CE",
    "voltage",
    "current",
    "feature_family",
    "cell_scope",
    "source_data_audit_only",
    "model_training_allowed",
]

WARNING_COLUMNS = [
    "workbook_name",
    "paper_id",
    "sheet_name",
    "warning_type",
    "warning_message",
    "source_data_audit_only",
    "model_training_allowed",
]

DEFAULT_SELECTED_SHEETS = {
    "41467_2025_63303_MOESM3_ESM.xlsx": [
        "Figure 3a",
        "Figure 3b",
        "Figure S7",
        "Figure S16",
    ],
    "41467_2025_66271_MOESM3_ESM.xlsx": [
        "Fig. 1f",
        "Fig. 1i",
        "Fig. 5c",
        "Fig. 6c",
        "Supplementary Fig. 36a",
        "Supplementary Fig. 37c",
        "Supplementary Fig. 42a",
        "Supplementary Fig. 43e",
    ],
}


def write_csv(path: Path, rows: list[dict[str, Any]], fieldnames: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def write_json(path: Path, payload: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def as_text(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def as_float(value: object) -> float | str:
    if value is None or value == "":
        return ""
    try:
        return float(value)
    except (TypeError, ValueError):
        return ""


def compact_float(value: object) -> str:
    parsed = as_float(value)
    if parsed == "":
        return ""
    return f"{parsed:.12g}"


def paper_id_for_workbook(path: Path) -> str:
    name = path.name
    if "63303" in name:
        return "ma_amanchukwu_2025_anode_free_active_learning"
    if "66271" in name:
        return "liu_chen_2025_initially_anode_free_charging_protocol"
    return "unknown_public_lmb_source"


def infer_cell_scope(workbook_name: str, sheet_name: str, sheet_type: str) -> str:
    text = f"{workbook_name} {sheet_name}".lower()
    if "63303" in workbook_name:
        if sheet_name in {"Figure S7"}:
            return "lmb_mechanism_test_not_full_cell"
        return "anode_free_full_cell"
    if "66271" in workbook_name:
        if sheet_name in {"Fig. 1f", "Fig. 1i"}:
            return "lmb_mechanism_test_not_full_cell"
        if "fig. 1" in text and sheet_type in {"time_voltage_current", "capacity_voltage_curve"}:
            return "lmb_mechanism_test_not_full_cell"
        return "anode_free_full_cell"
    return "unknown_scope"


def audit_use_for_type(sheet_type: str, cell_scope: str) -> str:
    if sheet_type == "cycle_capacity_ce":
        return "feature_schema_test;label_audit_candidate;not_training"
    if sheet_type == "time_voltage_current":
        return "time_voltage_feature_audit;strategy_feature_design;not_training"
    if sheet_type == "capacity_voltage_curve":
        return "curve_shape_feature_audit;not_training"
    if sheet_type == "electrolyte_descriptor":
        return "electrolyte_descriptor_audit;C20_proxy_only;not_full_RUL"
    if sheet_type == "active_learning_metric":
        return "strategy_optimization_reference;active_learning_metric_audit"
    if sheet_type == "spectroscopy_or_characterization":
        return "mechanism_context_only;not_training"
    return "audit_only_until_manual_review"


def preview_rows(ws: Any, max_preview_rows: int) -> list[list[str]]:
    rows: list[list[str]] = []
    max_rows = min(ws.max_row or 0, max_preview_rows)
    max_cols = min(ws.max_column or 0, 12)
    for row in ws.iter_rows(min_row=1, max_row=max_rows, max_col=max_cols, values_only=True):
        rows.append([as_text(value) for value in row])
    return rows


def flatten_preview(rows: list[list[str]]) -> str:
    compact = []
    for row in rows[:4]:
        nonempty = [value for value in row if value]
        if nonempty:
            compact.append(" | ".join(nonempty[:8]))
    return " || ".join(compact)


def infer_sheet_type(sheet_name: str, preview: list[list[str]]) -> str:
    text = f"{sheet_name} " + " ".join(value.lower() for row in preview for value in row if value)
    normalized_sheet_name = sheet_name.strip().lower()
    has_cycle = "cycle" in text
    has_capacity = "capacity" in text or re.search(r"\bcap\b", text) is not None
    has_ce = " ce" in f" {text}" or "coulomb" in text
    has_time = "time" in text
    has_voltage = "voltage" in text or "potential" in text or "ewe" in text
    has_current = "current" in text
    has_spectra = any(token in text for token in ["raman", "spectrum", "spectra", "peak", "energy", "xps", "nmr"])
    has_electrolyte = any(token in text for token in ["eli", "solvent", "substructure", "functionality", "electrolyte", "batch"])
    has_active_learning = any(token in text for token in ["fold", "matern", "pairwise", "shap", "feature", "mean_norm_capacity"])

    if normalized_sheet_name == "figure s16" or (has_electrolyte and "20th cycle" in text):
        return "electrolyte_descriptor"
    if has_cycle and (has_capacity or has_ce):
        return "cycle_capacity_ce"
    if has_time and has_voltage and has_current:
        return "time_voltage_current"
    if has_time and has_voltage:
        return "time_voltage_current"
    if has_capacity and has_voltage:
        return "capacity_voltage_curve"
    if has_electrolyte and (has_capacity or "eli" in text):
        return "electrolyte_descriptor"
    if has_active_learning:
        return "active_learning_metric"
    if has_spectra:
        return "spectroscopy_or_characterization"
    return "figure_source_unknown"


def load_sheet_values(ws: Any) -> list[list[Any]]:
    return [list(row) for row in ws.iter_rows(values_only=True)]


def row_has_token(row: list[Any], tokens: tuple[str, ...]) -> bool:
    text = " ".join(as_text(value).lower() for value in row)
    return any(token in text for token in tokens)


def find_cycle_header(values: list[list[Any]]) -> int | None:
    for index, row in enumerate(values[:8]):
        if row_has_token(row, ("cycle",)):
            return index
    return None


def find_time_header(values: list[list[Any]]) -> int | None:
    for index, row in enumerate(values[:8]):
        if row_has_token(row, ("time",)) and row_has_token(row, ("voltage", "potential", "current")):
            return index
    return None


def make_trace_name(parts: list[str]) -> str:
    clean = [part for part in parts if part and part.lower() not in {"capacity", "cap", "ce", "voltage", "potential", "current", "time"}]
    return "_".join(clean).strip("_") or "trace"


def split_measurement_header(raw: str) -> tuple[str, str]:
    text = raw.strip()
    lower = text.lower()
    if lower in {"capacity", "cap"} or "capacity" in lower or lower.endswith("-cap"):
        return text.rsplit("-", 1)[0] if "-" in text else "trace", "capacity"
    if lower in {"ce", "coulombic efficiency"} or lower.endswith("-ce"):
        return text.rsplit("-", 1)[0] if "-" in text else "trace", "CE"
    if "normalized" in lower and "capacity" in lower:
        return text, "normalized_capacity"
    if "voltage" in lower or "potential" in lower or "ewe" in lower:
        return text, "voltage"
    if "current" in lower:
        return text, "current"
    if "time" in lower:
        return text, "time"
    return text, "value"


def parse_cycle_capacity_ce(
    values: list[list[Any]],
    workbook_name: str,
    paper_id: str,
    sheet_name: str,
    cell_scope: str,
    max_rows: int,
) -> tuple[list[dict[str, Any]], list[str]]:
    header_index = find_cycle_header(values)
    if header_index is None:
        return [], ["No cycle header found"]

    header = values[header_index]
    group_header = values[header_index - 1] if header_index > 0 else []
    cycle_col = next((idx for idx, value in enumerate(header) if "cycle" in as_text(value).lower()), None)
    if cycle_col is None:
        return [], ["No cycle column found"]

    trace_measure_columns: dict[str, dict[str, int]] = defaultdict(dict)
    warnings: list[str] = []
    for col_index, raw_header in enumerate(header):
        if col_index == cycle_col:
            continue
        header_text = as_text(raw_header)
        parent_text = as_text(group_header[col_index]) if col_index < len(group_header) else ""
        if not header_text and not parent_text:
            continue
        if header_text.lower() in {"capacity", "ce"} and parent_text:
            trace = parent_text
            measurement = "capacity" if header_text.lower() == "capacity" else "CE"
        elif header_text.lower() in {"cell 1", "cell 2", "cell 3", "cell 4"}:
            trace = header_text
            measurement = "CE"
        elif header_text:
            trace, measurement = split_measurement_header(header_text)
            if measurement == "value":
                if parent_text and "capacity" in parent_text.lower():
                    measurement = "capacity"
                elif parent_text and "ce" in parent_text.lower():
                    measurement = "CE"
                else:
                    measurement = "capacity"
        else:
            continue
        trace_measure_columns[trace][measurement] = col_index

    if not trace_measure_columns:
        return [], ["No trace columns found"]

    long_rows: list[dict[str, Any]] = []
    for row in values[header_index + 1 : header_index + 1 + max_rows]:
        cycle_value = row[cycle_col] if cycle_col < len(row) else ""
        cycle = compact_float(cycle_value)
        if cycle == "":
            continue
        for trace, columns in trace_measure_columns.items():
            capacity = compact_float(row[columns["capacity"]]) if "capacity" in columns and columns["capacity"] < len(row) else ""
            normalized = ""
            ce_value = compact_float(row[columns["CE"]]) if "CE" in columns and columns["CE"] < len(row) else ""
            if "normalized" in trace.lower() and capacity:
                normalized = capacity
                capacity = ""
            if not any([capacity, normalized, ce_value]):
                continue
            long_rows.append(
                {
                    "workbook_name": workbook_name,
                    "paper_id": paper_id,
                    "sheet_name": sheet_name,
                    "trace_id": f"{sheet_name}:{trace}",
                    "inferred_cell_or_condition": trace,
                    "cycle_index": cycle,
                    "time_index": "",
                    "capacity": capacity,
                    "normalized_capacity": normalized,
                    "CE": ce_value,
                    "voltage": "",
                    "current": "",
                    "feature_family": "cycle_capacity_ce",
                    "cell_scope": cell_scope,
                    "source_data_audit_only": True,
                    "model_training_allowed": False,
                }
            )
    if not long_rows:
        warnings.append("No long-format cycle rows emitted")
    return long_rows, warnings


def parse_time_voltage_current(
    values: list[list[Any]],
    workbook_name: str,
    paper_id: str,
    sheet_name: str,
    cell_scope: str,
    max_rows: int,
) -> tuple[list[dict[str, Any]], list[str]]:
    header_index = find_time_header(values)
    if header_index is None:
        return [], ["No time-voltage/current header found"]

    header = values[header_index]
    group_header = values[header_index - 1] if header_index > 0 else []
    long_rows: list[dict[str, Any]] = []
    warnings: list[str] = []
    col = 0
    while col < len(header):
        header_text = as_text(header[col]).lower()
        if "time" not in header_text:
            col += 1
            continue
        trace = as_text(group_header[col]) or as_text(group_header[col + 1]) if col + 1 < len(group_header) else ""
        trace = trace or f"trace_{col}"
        time_col = col
        voltage_col = None
        current_col = None
        for maybe_col in range(col + 1, min(len(header), col + 4)):
            maybe = as_text(header[maybe_col]).lower()
            if voltage_col is None and ("voltage" in maybe or "potential" in maybe):
                voltage_col = maybe_col
            if current_col is None and "current" in maybe:
                current_col = maybe_col
        if voltage_col is None and current_col is None:
            col += 1
            continue
        for row in values[header_index + 1 : header_index + 1 + max_rows]:
            time_value = compact_float(row[time_col]) if time_col < len(row) else ""
            voltage = compact_float(row[voltage_col]) if voltage_col is not None and voltage_col < len(row) else ""
            current = compact_float(row[current_col]) if current_col is not None and current_col < len(row) else ""
            if not any([time_value, voltage, current]):
                continue
            long_rows.append(
                {
                    "workbook_name": workbook_name,
                    "paper_id": paper_id,
                    "sheet_name": sheet_name,
                    "trace_id": f"{sheet_name}:{trace}",
                    "inferred_cell_or_condition": trace,
                    "cycle_index": "",
                    "time_index": time_value,
                    "capacity": "",
                    "normalized_capacity": "",
                    "CE": "",
                    "voltage": voltage,
                    "current": current,
                    "feature_family": "time_voltage_current",
                    "cell_scope": cell_scope,
                    "source_data_audit_only": True,
                    "model_training_allowed": False,
                }
            )
        col += 2
    if not long_rows:
        warnings.append("No long-format time-voltage/current rows emitted")
    return long_rows, warnings


def parse_electrolyte_descriptor(
    values: list[list[Any]],
    workbook_name: str,
    paper_id: str,
    sheet_name: str,
    cell_scope: str,
    max_rows: int,
) -> tuple[list[dict[str, Any]], list[str]]:
    header_index = None
    for idx, row in enumerate(values[:6]):
        text = " ".join(as_text(value).lower() for value in row)
        if "eli" in text or ("discharge" in text and "capacity" in text):
            header_index = idx
            break
    if header_index is None:
        return [], ["No electrolyte descriptor header found"]

    header = [as_text(value) for value in values[header_index]]
    long_rows: list[dict[str, Any]] = []
    for row in values[header_index + 1 : header_index + 1 + max_rows]:
        condition = as_text(row[0]) if row else ""
        if not condition:
            continue
        row_map = {header[col]: row[col] for col in range(min(len(header), len(row))) if header[col]}
        capacity = ""
        for key, value in row_map.items():
            if "discharge" in key.lower() and "capacity" in key.lower():
                capacity = compact_float(value)
        long_rows.append(
            {
                "workbook_name": workbook_name,
                "paper_id": paper_id,
                "sheet_name": sheet_name,
                "trace_id": f"{sheet_name}:{condition}",
                "inferred_cell_or_condition": condition,
                "cycle_index": "",
                "time_index": "",
                "capacity": capacity,
                "normalized_capacity": "",
                "CE": "",
                "voltage": "",
                "current": "",
                "feature_family": "electrolyte_descriptor_C20_proxy_not_RUL",
                "cell_scope": cell_scope,
                "source_data_audit_only": True,
                "model_training_allowed": False,
            }
        )
    return long_rows, [] if long_rows else ["No electrolyte descriptor rows emitted"]


def parse_selected_sheet(
    values: list[list[Any]],
    workbook_name: str,
    paper_id: str,
    sheet_name: str,
    sheet_type: str,
    cell_scope: str,
    max_rows: int,
) -> tuple[list[dict[str, Any]], list[str]]:
    if sheet_type == "cycle_capacity_ce":
        return parse_cycle_capacity_ce(values, workbook_name, paper_id, sheet_name, cell_scope, max_rows)
    if sheet_type == "time_voltage_current":
        return parse_time_voltage_current(values, workbook_name, paper_id, sheet_name, cell_scope, max_rows)
    if sheet_type == "electrolyte_descriptor":
        return parse_electrolyte_descriptor(values, workbook_name, paper_id, sheet_name, cell_scope, max_rows)
    if sheet_type == "capacity_voltage_curve":
        return [], ["capacity_voltage_curve is inventoried but not long-parsed in v1 tiny validation"]
    return [], [f"{sheet_type} is audit-only or unsupported for long parsing in v1"]


def parse_selected_sheet_tokens(tokens: list[str] | None) -> dict[str, set[str]]:
    if not tokens:
        return {workbook: set(sheets) for workbook, sheets in DEFAULT_SELECTED_SHEETS.items()}
    selected: dict[str, set[str]] = defaultdict(set)
    for token in tokens:
        if "::" in token:
            workbook, sheet = token.split("::", 1)
            selected[workbook].add(sheet)
        else:
            selected["*"].add(token)
            for workbook, defaults in DEFAULT_SELECTED_SHEETS.items():
                if token in defaults:
                    selected[workbook].add(token)
    return selected


def build_report_md(report: dict[str, Any]) -> str:
    type_lines = "\n".join(
        f"- `{sheet_type}`: {count}" for sheet_type, count in sorted(report["sheet_type_distribution"].items())
    )
    parsed_lines = "\n".join(
        f"- {row['workbook_name']} / {row['sheet_name']}: {row['parse_status']} ({row['parsed_rows']} rows)"
        for row in report["selected_sheet_parse_manifest"]
    )
    warning_lines = "\n".join(
        f"- {row['workbook_name']} / {row['sheet_name']}: {row['warning_message']}"
        for row in report["warnings"]
    ) or "- No parser warnings."
    return f"""# Public LMB Nature Source Data Parser Audit

Generated: {report['generated_at_utc']}

This report is source-data audit only. Nature source-data workbooks are figure-source data, not standardized raw full-cell training datasets.

```text
source_data_audit_only=True
model_training_allowed=False
formal_model_performance_claimed=False
```

## Workbook Summary

| workbook | paper_id | sheets |
| --- | --- | --- |
{chr(10).join(f"| {row['workbook_name']} | {row['paper_id']} | {row['sheet_count']} |" for row in report['source_workbook_inventory'])}

## Sheet Type Distribution

{type_lines}

## Selected Sheet Tiny Parse

{parsed_lines}

## Parser Warnings

{warning_lines}

## Feature Schema Test Candidates

{chr(10).join(f"- {item}" for item in report['feature_schema_test_candidates'])}

## Audit-Only Sheets

{chr(10).join(f"- {item}" for item in report['audit_only_sheets'][:30])}

## Strategy Optimization Value

The Liu / Chen workbook is especially useful for strategy feature design because it contains CC vs MPC current, capacity, CE, and voltage traces. The Ma / Amanchukwu workbook is especially useful for electrolyte-screening reward proxies and active-learning design.

## Missing Metadata Before Training

- unified cell_id / trace_id manifest
- full cell versus mechanism-test scope review per trace
- cathode, anode/anode-free status, electrolyte, separator, pressure, temperature
- current density, areal capacity, voltage limits, planned cycle count
- termination reason and failure mode
- observed / censored / protocol_censored label audit

## Gate Decision

```text
public_source_feature_schema_test_allowed=True
public_source_label_audit_allowed=True
direct_full_cell_training_allowed=False
model_training_allowed=False
```

Next allowed step: write a public source feature-schema test, still audit-only.
"""


def parse_public_lmb_nature_source_data(
    input_files: list[Path],
    output_root: Path,
    overwrite: bool = False,
    max_preview_rows: int = 100,
    selected_sheets: list[str] | None = None,
) -> dict[str, Any]:
    if output_root.exists() and any(output_root.iterdir()) and not overwrite:
        raise FileExistsError(f"Output root already exists and is not empty: {output_root}")
    output_root.mkdir(parents=True, exist_ok=True)

    selected = parse_selected_sheet_tokens(selected_sheets)
    workbook_rows: list[dict[str, Any]] = []
    sheet_rows: list[dict[str, Any]] = []
    manifest_rows: list[dict[str, Any]] = []
    long_rows: list[dict[str, Any]] = []
    warning_rows: list[dict[str, Any]] = []
    feature_schema_candidates: list[str] = []
    audit_only_sheets: list[str] = []

    for input_file in input_files:
        if not input_file.exists():
            raise FileNotFoundError(input_file)
        workbook_name = input_file.name
        paper_id = paper_id_for_workbook(input_file)
        workbook = openpyxl.load_workbook(input_file, read_only=True, data_only=True)
        try:
            workbook_rows.append(
                {
                    "workbook_name": workbook_name,
                    "source_path": str(input_file),
                    "paper_id": paper_id,
                    "sheet_count": len(workbook.sheetnames),
                    "source_data_audit_only": True,
                    "model_training_allowed": False,
                }
            )
            workbook_selected = set(selected.get(workbook_name, set())) | set(selected.get("*", set()))
            for worksheet in workbook.worksheets:
                preview = preview_rows(worksheet, min(max_preview_rows, 8))
                inferred_type = infer_sheet_type(worksheet.title, preview)
                cell_scope = infer_cell_scope(workbook_name, worksheet.title, inferred_type)
                audit_use = audit_use_for_type(inferred_type, cell_scope)
                sheet_row = {
                    "workbook_name": workbook_name,
                    "paper_id": paper_id,
                    "sheet_name": worksheet.title,
                    "n_rows": worksheet.max_row,
                    "n_cols": worksheet.max_column,
                    "header_preview": flatten_preview(preview),
                    "inferred_sheet_type": inferred_type,
                    "dataset_role": "true_lmb",
                    "cell_scope": cell_scope,
                    "training_allowed_now": False,
                    "audit_use": audit_use,
                    "source_data_audit_only": True,
                    "model_training_allowed": False,
                }
                sheet_rows.append(sheet_row)
                sheet_label = f"{workbook_name}::{worksheet.title}"
                if inferred_type in {"cycle_capacity_ce", "time_voltage_current", "capacity_voltage_curve", "electrolyte_descriptor"}:
                    feature_schema_candidates.append(sheet_label)
                if "not_training" in audit_use or inferred_type in {"spectroscopy_or_characterization", "active_learning_metric", "figure_source_unknown"}:
                    audit_only_sheets.append(sheet_label)

                if worksheet.title not in workbook_selected:
                    continue

                values = load_sheet_values(worksheet)
                parsed_rows, warnings = parse_selected_sheet(
                    values=values,
                    workbook_name=workbook_name,
                    paper_id=paper_id,
                    sheet_name=worksheet.title,
                    sheet_type=inferred_type,
                    cell_scope=cell_scope,
                    max_rows=max_preview_rows,
                )
                long_rows.extend(parsed_rows)
                for warning in warnings:
                    warning_rows.append(
                        {
                            "workbook_name": workbook_name,
                            "paper_id": paper_id,
                            "sheet_name": worksheet.title,
                            "warning_type": "selected_sheet_parse_warning",
                            "warning_message": warning,
                            "source_data_audit_only": True,
                            "model_training_allowed": False,
                        }
                    )
                manifest_rows.append(
                    {
                        "workbook_name": workbook_name,
                        "paper_id": paper_id,
                        "sheet_name": worksheet.title,
                        "inferred_sheet_type": inferred_type,
                        "parse_status": "parsed_with_warnings" if warnings and parsed_rows else "parsed" if parsed_rows else "warning_only",
                        "parsed_rows": len(parsed_rows),
                        "parser_warning": ";".join(warnings),
                        "source_data_audit_only": True,
                        "model_training_allowed": False,
                    }
                )
        finally:
            workbook.close()

    selected_missing = []
    available = {(row["workbook_name"], row["sheet_name"]) for row in sheet_rows}
    for workbook_name, sheets in selected.items():
        for sheet_name in sheets:
            if (workbook_name, sheet_name) not in available:
                selected_missing.append(f"{workbook_name}::{sheet_name}")
                warning_rows.append(
                    {
                        "workbook_name": workbook_name,
                        "paper_id": paper_id_for_workbook(Path(workbook_name)),
                        "sheet_name": sheet_name,
                        "warning_type": "selected_sheet_missing",
                        "warning_message": "Selected sheet was not found in the workbook",
                        "source_data_audit_only": True,
                        "model_training_allowed": False,
                    }
                )

    type_distribution = dict(Counter(row["inferred_sheet_type"] for row in sheet_rows))
    report = {
        "generated_at_utc": datetime.now(timezone.utc).isoformat(),
        "source_data_audit_only": True,
        "model_training_allowed": False,
        "formal_model_performance_claimed": False,
        "source_workbook_inventory": workbook_rows,
        "sheet_count_total": len(sheet_rows),
        "sheet_type_distribution": type_distribution,
        "selected_sheet_parse_manifest": manifest_rows,
        "long_preview_rows": len(long_rows),
        "warnings": warning_rows,
        "selected_missing": selected_missing,
        "feature_schema_test_candidates": feature_schema_candidates,
        "audit_only_sheets": audit_only_sheets,
        "public_source_feature_schema_test_allowed": True,
        "direct_full_cell_training_allowed": False,
    }

    write_csv(output_root / "source_workbook_inventory.csv", workbook_rows, WORKBOOK_INVENTORY_COLUMNS)
    write_csv(output_root / "source_sheet_inventory.csv", sheet_rows, SHEET_INVENTORY_COLUMNS)
    write_csv(output_root / "selected_sheet_parse_manifest.csv", manifest_rows, PARSE_MANIFEST_COLUMNS)
    write_csv(output_root / "public_lmb_source_long_preview.csv", long_rows, LONG_COLUMNS)
    write_csv(output_root / "source_data_parser_warnings.csv", warning_rows, WARNING_COLUMNS)
    write_json(output_root / "public_lmb_source_data_audit_report.json", report)
    (output_root / "public_lmb_source_data_audit_report.md").write_text(build_report_md(report), encoding="utf-8")
    return report


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input-files", nargs="+", required=True, type=Path)
    parser.add_argument("--output-root", required=True, type=Path)
    parser.add_argument("--overwrite", action="store_true")
    parser.add_argument("--max-preview-rows", type=int, default=100)
    parser.add_argument("--selected-sheets", nargs="*")
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    parse_public_lmb_nature_source_data(
        input_files=args.input_files,
        output_root=args.output_root,
        overwrite=args.overwrite,
        max_preview_rows=args.max_preview_rows,
        selected_sheets=args.selected_sheets,
    )


if __name__ == "__main__":
    main()
