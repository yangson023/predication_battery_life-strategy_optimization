"""Validate partner-filled lithium metal battery metadata workbooks.

This validator is a metadata gate only. It does not read BTSDA raw data, create
labels, train models, or enter the RUL prediction pipeline.
"""

from __future__ import annotations

import argparse
import csv
import json
import sys
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


SHEET_NAME = "填写模板"

P0_FIELDS = [
    "电池编号",
    "电池类型",
    "原始数据文件夹",
    "实验日期",
    "正负极体系",
    "电解液溶剂",
    "锂盐",
    "锂盐浓度",
    "电解液添加剂",
    "隔膜材料/型号",
    "电流密度(mA/cm²)",
    "面容量(mAh/cm²)",
    "循环协议",
    "截止条件",
    "终止原因",
    "失效模式/现象",
]

P1_FIELDS = [
    "实验批次",
    "操作者",
    "测试设备/通道",
    "负极/基底材料",
    "锂片厚度(μm)",
    "铜箔/基底信息",
    "电池壳类型",
    "电解液用量(μL)",
    "电极面积(cm²)",
    "压力条件",
    "测试温度(℃)",
    "化成协议",
    "静置时间",
    "是否协议变化",
    "是否异常实验",
    "保密/公开级别",
]

CRITICAL_LMB_FIELDS = [
    "电解液溶剂",
    "锂盐",
    "锂盐浓度",
    "电解液添加剂",
    "隔膜材料/型号",
    "电流密度(mA/cm²)",
    "面容量(mAh/cm²)",
]

VALID_CELL_TYPES = {
    "Li||Li 对称电池": "Li||Li",
    "Li||Cu 半电池": "Li||Cu",
    "Li||全电池": "Li||Full",
    "Anode-free 全电池": "Anode-free",
    "其他": "Other",
    "未知": "Unknown",
}

TERMINATION_CATEGORY = {
    "自然失效": "observed_candidate",
    "短路": "observed_candidate",
    "人为停止": "protocol_censored",
    "设备中断": "protocol_censored",
    "协议结束": "protocol_censored",
    "数据导出不完整": "protocol_censored",
    "未记录": "unknown",
}

PER_CELL_COLUMNS = [
    "电池编号",
    "cell_group",
    "电池类型",
    "原始数据文件夹",
    "metadata_gate_status",
    "missing_p0_count",
    "missing_p1_count",
    "missing_critical_lmb_fields",
    "termination_interpretation",
    "protocol_change_filled",
    "abnormal_run_filled",
    "confidentiality_filled",
    "label_audit_allowed",
    "trainability_audit_prerequisite",
    "training_allowed_now",
    "gate_reason",
]

MISSING_FIELD_COLUMNS = [
    "电池编号",
    "cell_group",
    "field_priority",
    "field_name",
    "missing_reason",
]

SUMMARY_COLUMNS = [
    "metric",
    "value",
]


def normalize_cell(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    return "" if text.lower() == "none" else text


def is_blank(value: Any) -> bool:
    return normalize_cell(value) == ""


def read_metadata_workbook(path: Path) -> list[dict[str, str]]:
    workbook = load_workbook(path, data_only=True, read_only=True)
    try:
        if SHEET_NAME not in workbook.sheetnames:
            raise ValueError(f"Workbook must contain sheet: {SHEET_NAME}")
        sheet = workbook[SHEET_NAME]
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [normalize_cell(value) for value in rows[0]]
        result: list[dict[str, str]] = []
        for row in rows[1:]:
            record = {headers[index]: normalize_cell(value) for index, value in enumerate(row) if index < len(headers)}
            if not record.get("电池编号") and not record.get("电池类型") and not record.get("原始数据文件夹"):
                continue
            result.append(record)
        return result
    finally:
        workbook.close()


def field_missing(row: dict[str, str], field_name: str) -> bool:
    return is_blank(row.get(field_name))


def cell_group_for(cell_type: str) -> str:
    return VALID_CELL_TYPES.get(cell_type, "Unknown")


def termination_interpretation(reason: str) -> str:
    if not reason:
        return "unknown"
    normalized = reason.strip()
    if any(keyword in normalized for keyword in ["到达设定循环数", "达到设定循环数", "到达计划循环数", "达到计划循环数"]):
        return "protocol_censored"
    return TERMINATION_CATEGORY.get(reason, "needs_manual_review")


def gate_cell(row: dict[str, str]) -> tuple[dict[str, Any], list[dict[str, str]]]:
    cell_id = row.get("电池编号", "")
    cell_type = row.get("电池类型", "")
    cell_group = cell_group_for(cell_type)
    missing_p0 = [field for field in P0_FIELDS if field_missing(row, field)]
    missing_p1 = [field for field in P1_FIELDS if field_missing(row, field)]
    missing_critical = [field for field in CRITICAL_LMB_FIELDS if field_missing(row, field)]
    termination = termination_interpretation(row.get("终止原因", ""))
    unknown_cell_type = cell_group == "Unknown"

    if unknown_cell_type:
        status = "metadata_unknown_cell_type"
        reason = "电池类型未知或不在允许选项中。"
    elif missing_p0:
        status = "metadata_blocked_missing_p0"
        reason = "缺失 P0 关键字段，不能进入 label trainability audit。"
    elif missing_p1:
        status = "metadata_partial_needs_review"
        reason = "P0 完整但 P1 尚不完整，可做初步 label audit，但需要人工复核。"
    else:
        status = "metadata_ready_for_label_audit"
        reason = "P0/P1 字段完整，可进入 label audit 前置复核。"

    if termination in {"unknown", "needs_manual_review"} and status == "metadata_ready_for_label_audit":
        status = "metadata_partial_needs_review"
        reason = "终止原因仍需人工解释，不能直接判断 observed/censored。"

    missing_rows: list[dict[str, str]] = []
    for field in missing_p0:
        missing_rows.append(
            {
                "电池编号": cell_id,
                "cell_group": cell_group,
                "field_priority": "P0",
                "field_name": field,
                "missing_reason": "P0 必填字段缺失",
            }
        )
    for field in missing_p1:
        missing_rows.append(
            {
                "电池编号": cell_id,
                "cell_group": cell_group,
                "field_priority": "P1",
                "field_name": field,
                "missing_reason": "P1 推荐字段缺失",
            }
        )

    label_audit_allowed = status in {"metadata_ready_for_label_audit", "metadata_partial_needs_review"}
    gate_row = {
        "电池编号": cell_id,
        "cell_group": cell_group,
        "电池类型": cell_type,
        "原始数据文件夹": row.get("原始数据文件夹", ""),
        "metadata_gate_status": status,
        "missing_p0_count": len(missing_p0),
        "missing_p1_count": len(missing_p1),
        "missing_critical_lmb_fields": ";".join(missing_critical),
        "termination_interpretation": termination,
        "protocol_change_filled": not field_missing(row, "是否协议变化"),
        "abnormal_run_filled": not field_missing(row, "是否异常实验"),
        "confidentiality_filled": not field_missing(row, "保密/公开级别"),
        "label_audit_allowed": label_audit_allowed,
        "trainability_audit_prerequisite": status == "metadata_ready_for_label_audit",
        "training_allowed_now": False,
        "gate_reason": reason,
    }
    return gate_row, missing_rows


def write_csv(path: Path, rows: list[dict[str, Any]], fieldnames: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def build_summary(gate_rows: list[dict[str, Any]], missing_rows: list[dict[str, str]]) -> list[dict[str, Any]]:
    status_counts = Counter(row["metadata_gate_status"] for row in gate_rows)
    group_counts = Counter(row["cell_group"] for row in gate_rows)
    p0_missing_count = sum(1 for row in missing_rows if row["field_priority"] == "P0")
    p1_missing_count = sum(1 for row in missing_rows if row["field_priority"] == "P1")
    summary = [
        {"metric": "cell_count", "value": len(gate_rows)},
        {"metric": "training_allowed_now", "value": False},
        {"metric": "label_audit_ready_cell_count", "value": status_counts.get("metadata_ready_for_label_audit", 0)},
        {"metric": "metadata_partial_needs_review_count", "value": status_counts.get("metadata_partial_needs_review", 0)},
        {"metric": "metadata_blocked_missing_p0_count", "value": status_counts.get("metadata_blocked_missing_p0", 0)},
        {"metric": "metadata_unknown_cell_type_count", "value": status_counts.get("metadata_unknown_cell_type", 0)},
        {"metric": "total_missing_p0_fields", "value": p0_missing_count},
        {"metric": "total_missing_p1_fields", "value": p1_missing_count},
    ]
    for group, count in sorted(group_counts.items()):
        summary.append({"metric": f"cell_group_count::{group}", "value": count})
    return summary


def write_report(
    output_root: Path,
    workbook_path: Path,
    gate_rows: list[dict[str, Any]],
    missing_rows: list[dict[str, str]],
    summary_rows: list[dict[str, Any]],
) -> dict[str, Any]:
    status_counts = Counter(row["metadata_gate_status"] for row in gate_rows)
    group_status_counts: dict[str, Counter[str]] = defaultdict(Counter)
    for row in gate_rows:
        group_status_counts[str(row["cell_group"])][str(row["metadata_gate_status"])] += 1
    missing_by_field = Counter(row["field_name"] for row in missing_rows)
    report = {
        "input_workbook": str(workbook_path),
        "sheet_name": SHEET_NAME,
        "training_allowed_now": False,
        "model_training_allowed": False,
        "labels_generated": False,
        "raw_btsda_data_read": False,
        "cell_count": len(gate_rows),
        "metadata_gate_status_counts": dict(status_counts),
        "group_status_counts": {group: dict(counter) for group, counter in group_status_counts.items()},
        "top_missing_fields": dict(missing_by_field.most_common(20)),
        "summary": summary_rows,
        "notes": [
            "Metadata validation is a pre-label-audit gate, not model performance.",
            "Partner handwritten notes must be copied into the structured Excel fields before they can pass this gate.",
            "Li||Li and Li||Cu are counted separately.",
        ],
    }
    with (output_root / "lmb_metadata_validation_report.json").open("w", encoding="utf-8") as handle:
        json.dump(report, handle, ensure_ascii=False, indent=2)

    markdown = [
        "# LMB Metadata Validation Report",
        "",
        "This report validates partner-filled metadata only. It does not read raw BTSDA data, create labels, or train models.",
        "",
        f"- Input workbook: `{workbook_path}`",
        f"- Sheet: `{SHEET_NAME}`",
        f"- Cell count: {len(gate_rows)}",
        f"- Training allowed now: {report['training_allowed_now']}",
        "",
        "## Gate Status Counts",
        "",
        "| status | count |",
        "| --- | ---: |",
    ]
    for status, count in sorted(status_counts.items()):
        markdown.append(f"| `{status}` | {count} |")
    markdown.extend(["", "## Top Missing Fields", "", "| field | missing count |", "| --- | ---: |"])
    for field, count in missing_by_field.most_common(20):
        markdown.append(f"| `{field}` | {count} |")
    markdown.extend(
        [
            "",
            "## Gate Decision",
            "",
            "- `training_allowed_now = False`",
            "- `model_training_allowed = False`",
            "- Current metadata must be completed before label trainability audit.",
        ]
    )
    (output_root / "lmb_metadata_validation_report.md").write_text("\n".join(markdown) + "\n", encoding="utf-8")
    return report


def validate_metadata_template(input_workbook: Path, output_root: Path) -> dict[str, Any]:
    rows = read_metadata_workbook(input_workbook)
    gate_rows: list[dict[str, Any]] = []
    missing_rows: list[dict[str, str]] = []
    for row in rows:
        gate_row, cell_missing_rows = gate_cell(row)
        gate_rows.append(gate_row)
        missing_rows.extend(cell_missing_rows)

    summary_rows = build_summary(gate_rows, missing_rows)
    output_root.mkdir(parents=True, exist_ok=True)
    write_csv(output_root / "per_cell_metadata_gate.csv", gate_rows, PER_CELL_COLUMNS)
    write_csv(output_root / "missing_field_report.csv", missing_rows, MISSING_FIELD_COLUMNS)
    write_csv(output_root / "metadata_validation_summary.csv", summary_rows, SUMMARY_COLUMNS)
    return write_report(output_root, input_workbook, gate_rows, missing_rows, summary_rows)


def build_arg_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input-workbook", required=True, help="Path to partner-filled metadata xlsx.")
    parser.add_argument("--output-root", required=True, help="Directory for validation outputs.")
    return parser


def main() -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    args = build_arg_parser().parse_args()
    report = validate_metadata_template(Path(args.input_workbook), Path(args.output_root))
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
